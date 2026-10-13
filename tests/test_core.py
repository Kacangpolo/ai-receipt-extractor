"""Offline tests: no API key, no database needed.  Run: python -m pytest -q"""
import csv
import json
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

from src.evaluate import evaluate
from src.extractor import ReceiptExtractor, build_request, parse_response
from src.pipeline import run
from src.schema import Receipt

SAMPLE = {
    "merchant_name": "Warung Bu Made",
    "receipt_date": "2026-10-10",
    "items": [
        {"name": "Nasi Campur", "quantity": 2, "unit_price": 20000, "line_total": 40000},
        {"name": "Es Teh", "quantity": 2, "unit_price": 5000, "line_total": 10000},
    ],
    "tax": 0,
    "total": 50000,
    "confidence": 0.95,
}

# Smallest valid PNG, so the pipeline has a real image file to read.
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082"
)


def fake_response(payload):
    block = SimpleNamespace(type="tool_use", name="record_receipt", input=payload)
    return SimpleNamespace(content=[block], usage=SimpleNamespace(input_tokens=1500, output_tokens=200))


class FakeClient:
    def __init__(self, payload):
        self.messages = SimpleNamespace(create=lambda **kw: fake_response(payload))


def test_parse_and_sum_check():
    r = parse_response(fake_response(SAMPLE))
    assert r.merchant_name == "Warung Bu Made"
    assert r.total == Decimal("50000")
    assert r.total_matches_items() is True


def test_sum_check_flags_mismatch():
    bad = dict(SAMPLE, total=90000)
    assert Receipt.model_validate(bad).total_matches_items() is False


def test_request_forces_tool_and_embeds_image(tmp_path):
    img = tmp_path / "n.png"
    img.write_bytes(PNG)
    req = build_request(img, model="test-model")
    assert req["tool_choice"] == {"type": "tool", "name": "record_receipt"}
    assert req["messages"][0]["content"][0]["source"]["media_type"] == "image/png"


def test_pipeline_and_evaluate_end_to_end(tmp_path):
    inp, out = tmp_path / "in", tmp_path / "out"
    inp.mkdir()
    (inp / "nota_01.png").write_bytes(PNG)
    ex = ReceiptExtractor(client=FakeClient(SAMPLE), model="test-model")

    stats = run(inp, out, use_db=False, limit=None, extractor=ex)
    assert stats["processed"] == 1 and stats["input_tokens"] == 1500
    assert json.loads((out / "nota_01.json").read_text())["total"] == "50000"

    # second run skips the already-processed image
    assert run(inp, out, use_db=False, limit=None, extractor=ex)["skipped"] == 1

    labels = tmp_path / "labels.csv"
    with labels.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["file", "merchant_name", "receipt_date", "total"])
        w.writerow(["nota_01.png", "Bu Made", "2026-10-10", "50000"])
    r = evaluate(labels, out)
    assert r["fully_correct_pct"] == 100.0


# ---- Gemini backend (fake clients, no network) ----------------------------
from src.gemini import GeminiExtractor, parse_json_text


class FakeInteractions:
    def __init__(self, payload, fail_times=0):
        self.payload, self.fail_times, self.calls = payload, fail_times, []

    def create(self, **kw):
        self.calls.append(kw)
        if len(self.calls) <= self.fail_times:
            raise RuntimeError("429 RESOURCE_EXHAUSTED")
        usage = SimpleNamespace(total_input_tokens=1200, total_output_tokens=150)
        return SimpleNamespace(output_text=json.dumps(self.payload), usage=usage)


def test_parse_json_text_strips_fences():
    r = parse_json_text("```json\n" + json.dumps(SAMPLE) + "\n```")
    assert r.total == Decimal("50000")


def test_gemini_extract_with_retry(tmp_path):
    img = tmp_path / "n.png"
    img.write_bytes(PNG)
    fake = FakeInteractions(SAMPLE, fail_times=1)
    ex = GeminiExtractor(client=SimpleNamespace(interactions=fake), model="m", min_interval=0)
    import src.gemini as g
    g._sleep = lambda s: None  # don't actually wait during tests
    receipt, tokens = ex.extract(img)
    assert receipt.merchant_name == "Warung Bu Made"
    assert tokens == {"input_tokens": 1200, "output_tokens": 150}
    assert len(fake.calls) == 2
    sent = fake.calls[-1]
    assert sent["response_format"]["mime_type"] == "application/json"
    assert sent["input"][1]["type"] == "image" and sent["input"][1]["mime_type"] == "image/png"


# ---- Date handling: model copies text, code parses (day-first) -------------
from src.schema import parse_printed_date
from datetime import date as _date


def test_printed_date_is_day_first():
    assert parse_printed_date("Tgl: 03/10/2026  Jam: 14:22") == _date(2026, 10, 3)
    assert parse_printed_date("07-10-26") == _date(2026, 10, 7)
    assert parse_printed_date("2026-10-10") == _date(2026, 10, 10)
    assert parse_printed_date("tidak terbaca") is None


def test_printed_date_overrides_model_guess():
    # The model swapped day/month (real bug seen with gemini on nota_01.jpg)
    r = Receipt.model_validate(dict(SAMPLE, receipt_date="2026-03-03", date_printed="03/10/2026"))
    assert r.receipt_date == _date(2026, 10, 3)


def test_amounts_drop_trailing_zero():
    assert str(Receipt.model_validate(dict(SAMPLE, total=149850.0)).total) == "149850"


# ---- Report / dashboard data ------------------------------------------------
from src.report import daily_spend, load_results, needs_review, to_excel_bytes, top_items, top_merchants


def test_report_tables_and_review(tmp_path):
    good = Receipt.model_validate(dict(SAMPLE, date_printed="10/10/2026"))
    bad = Receipt.model_validate(dict(SAMPLE, merchant_name="Toko X", total=99000,
                                      date_printed="11/10/2026", confidence=0.4))
    (tmp_path / "a.json").write_text(good.model_dump_json())
    (tmp_path / "b.json").write_text(bad.model_dump_json())
    (tmp_path / "junk.json").write_text('{"not": "a receipt", "items": "x"}')

    rec, itm = load_results(tmp_path)
    assert len(rec) == 2 and len(itm) == 4
    assert rec["total"].sum() == 149000

    rv = needs_review(rec)
    assert list(rv["file"]) == ["b"]
    assert "!= total" in rv.iloc[0]["reason"] and "low confidence" in rv.iloc[0]["reason"]

    assert len(daily_spend(rec)) == 2
    assert top_merchants(rec).iloc[0]["merchant_name"] == "Toko X"
    ti = top_items(itm)
    assert ti.iloc[0]["item"] == "Nasi Campur" and ti.iloc[0]["spend"] == 80000

    xlsx = to_excel_bytes(rec, itm)
    import openpyxl, io
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Receipts", "Items"] and wb["Items"].max_row == 5


class FlakyJSON(FakeInteractions):
    """First answer is truncated garbage (like nota_12), second is fine."""
    def create(self, **kw):
        self.calls.append(kw)
        usage = SimpleNamespace(total_input_tokens=1, total_output_tokens=1)
        text = '{"merchant_name": "X", "items": [{"name": "aaaa' if len(self.calls) == 1 else json.dumps(self.payload)
        return SimpleNamespace(output_text=text, usage=usage)


def test_gemini_retries_on_broken_json(tmp_path):
    img = tmp_path / "n.png"
    img.write_bytes(PNG)
    fake = FlakyJSON(SAMPLE)
    ex = GeminiExtractor(client=SimpleNamespace(interactions=fake), model="m", min_interval=0)
    receipt, _ = ex.extract(img)
    assert receipt.total == Decimal("50000") and len(fake.calls) == 2


# ---- Consistency checks + second opinion (real failure seen on nota_01, 13 Oct) ----
REAL_BAD_NOTA_01 = {  # what gemini-3.5-flash-lite actually returned: 1 of 5 items, total = Aqua's price
    "merchant_name": None, "date_printed": "03/10/2026", "currency": "IDR",
    "items": [{"name": "Aqua 600ml", "quantity": "3", "unit_price": "0", "line_total": None}],
    "subtotal": "85000", "tax": None, "total": "12000", "confidence": 1.0, "notes": None,
}
GOOD_NOTA_01 = {
    "merchant_name": "TOKO SARI MURNI", "date_printed": "03/10/2026",
    "items": [{"name": "Aqua 600ml", "quantity": 3, "unit_price": 4000},
              {"name": "Indomie Goreng", "quantity": 5, "unit_price": 3500},
              {"name": "Teh Pucuk 350ml", "quantity": 2, "unit_price": 4500},
              {"name": "Roti Tawar Sari", "quantity": 1, "unit_price": 17500},
              {"name": "Telur Ayam 1kg", "quantity": 1, "unit_price": 29000}],
    "subtotal": 85000, "tax": 9350, "total": 94350, "confidence": 0.9,
}


def test_real_failure_is_flagged():
    issues = Receipt.model_validate(REAL_BAD_NOTA_01).issues()
    assert "merchant missing" in issues
    assert "item without a price" in issues
    assert any("subtotal + tax" in i and "!= total" in i for i in issues)


def test_good_receipt_passes_and_line_totals_are_derived():
    r = Receipt.model_validate(GOOD_NOTA_01)
    assert r.items[0].line_total == Decimal("12000")  # 3 x 4000 computed by code
    assert r.issues() == []


def test_missing_item_breaks_subtotal_check():
    short = dict(GOOD_NOTA_01, items=GOOD_NOTA_01["items"][:1])
    assert any("!= subtotal" in i for i in Receipt.model_validate(short).issues())


class OneAnswer:
    def __init__(self, payload, model):
        self.payload, self.model, self.calls = payload, model, 0

    def extract(self, path):
        self.calls += 1
        return Receipt.model_validate(self.payload), {"input_tokens": 10, "output_tokens": 5}


def test_flagged_receipt_gets_second_opinion(tmp_path):
    inp, out = tmp_path / "in", tmp_path / "out"
    inp.mkdir()
    (inp / "nota_01.png").write_bytes(PNG)
    weak, strong = OneAnswer(REAL_BAD_NOTA_01, "small"), OneAnswer(GOOD_NOTA_01, "big")
    stats = run(inp, out, use_db=False, limit=None, extractor=weak, fallback=strong)
    saved = json.loads((out / "nota_01.json").read_text())
    assert stats["retried"] == 1 and stats["flagged"] == 0 and strong.calls == 1
    assert saved["_meta"] == {"model": "big", "issues": []} and saved["total"] == "94350"
    assert stats["input_tokens"] == 20  # both calls are counted


def test_good_receipt_skips_second_opinion(tmp_path):
    inp, out = tmp_path / "in", tmp_path / "out"
    inp.mkdir()
    (inp / "nota_01.png").write_bytes(PNG)
    strong = OneAnswer(GOOD_NOTA_01, "big")
    run(inp, out, use_db=False, limit=None, extractor=OneAnswer(GOOD_NOTA_01, "small"), fallback=strong)
    assert strong.calls == 0
